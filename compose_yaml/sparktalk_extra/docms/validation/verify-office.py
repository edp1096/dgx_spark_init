#!/usr/bin/env python3
"""Independent checks for the four published samples; not a general Office certifier."""
import datetime as dt
import base64
import hashlib
import html
import json
import math
import re
import subprocess
import sys
import tempfile
import xml.etree.ElementTree as ET
from pathlib import Path
from zipfile import ZipFile

from PIL import Image, __version__ as PIL_VERSION
import openpyxl

ROOT = Path(sys.argv[1])
INPUTS = {item['stem']: item for item in json.loads((ROOT / 'inputs.json').read_text())}
NS = {'w': 'http://schemas.openxmlformats.org/wordprocessingml/2006/main',
      'a': 'http://schemas.openxmlformats.org/drawingml/2006/main',
      'p': 'http://schemas.openxmlformats.org/presentationml/2006/main',
      'h': 'http://www.w3.org/1999/xhtml'}
RESULT = {'status': 'running', 'checks': [], 'pdfs': [], 'negative_controls': [], 'artifacts': [],
          'environment': json.loads((ROOT / 'environment.json').read_text())}

RESULT['controllers'] = {stem:json.loads((ROOT/stem/'reference/controller.json').read_text()) for stem in ('report','slides','sales')}
RESULT['environment'].update({'python':sys.version.split()[0],'openpyxl':openpyxl.__version__,'pillow':PIL_VERSION})

def require(condition, message):
    if not condition:
        raise AssertionError(message)

def normal(text):
    return re.sub(r'\s+', '', text)

def check(label, function):
    detail = function()
    RESULT['checks'].append({'name': label, 'status': 'pass', 'detail': detail})

def read_xml(archive, name):
    return ET.fromstring(archive.read(name))

def docx_snapshot(filename):
    with ZipFile(filename) as archive:
        require(archive.testzip() is None, 'DOCX ZIP CRC failure')
        root = read_xml(archive, 'word/document.xml')
        body = root.find('w:body', NS)
        paragraphs = []
        for p in body.findall('w:p', NS):
            value = ''.join(t.text or '' for t in p.findall('.//w:t', NS))
            if value:
                paragraphs.append(value)
        tables = [[[''.join(t.text or '' for t in c.findall('.//w:t', NS))
                    for c in row.findall('w:tc', NS)] for row in table.findall('w:tr', NS)]
                  for table in body.findall('w:tbl', NS)]
        return {'paragraphs': paragraphs, 'tables': tables}

def expected_docx():
    sample = INPUTS['report']
    return {'paragraphs': [sample['title'], sample['sections'][0]['heading'], *sample['sections'][0]['paragraphs']],
            'tables': [sample['sections'][0]['table']]}

def check_docx(filename):
    actual = docx_snapshot(filename)
    require(actual == expected_docx(), 'DOCX paragraph/table contents or order changed')
    return '3 paragraphs and all 81 × 2 table cells match input'

def pptx_snapshot(filename):
    with ZipFile(filename) as archive:
        require(archive.testzip() is None, 'PPTX ZIP CRC failure')
        names = sorted((name for name in archive.namelist() if re.fullmatch(r'ppt/slides/slide\d+\.xml', name)),
                       key=lambda value: int(re.search(r'(\d+)\.xml', value)[1]))
        slides = []
        for name in names:
            root = read_xml(archive, name)
            slides.append([''.join(t.text or '' for t in p.findall('.//a:t', NS))
                           for p in root.findall('.//a:p', NS) if p.findall('.//a:t', NS)])
        return slides

def expected_pptx():
    sample = INPUTS['slides']
    return [[slide['title'], *slide['bullets'], f'{i+1} / {len(sample["slides"])}']
            for i, slide in enumerate(sample['slides'])]

def check_pptx(filename):
    require(pptx_snapshot(filename) == expected_pptx(), 'PPTX slide items or order changed')
    return '2 slides: every title, bullet and page number matches input'

# Expected answers are independently specified, not read back from a generation engine.
ANSWERS = {('매출', 'D2'): 10, ('매출', 'D3'): 20, ('매출', 'C4'): 300, ('매출', 'D4'): 30,
           ('요약', 'B2'): 300, ('요약', 'B3'): 150, ('요약', 'B4'): True,
           ('요약', 'B5'): '충족', ('요약', 'B6'): '123', ('요약', 'B7'): ''}

def formula_normal(value):
    # LibreOffice omits optional quotes on these Korean sheet names and writes TRUE()/FALSE().
    value = re.sub(r"'(매출|요약)'!", r'\1!', value.lstrip('='))
    return re.sub(r'\b(TRUE|FALSE)\(\)', r'\1', value)

def same_value(actual, expected):
    if isinstance(expected, bool):
        return type(actual) is bool and actual == expected
    if isinstance(expected, (int, float)):
        return type(actual) in (int, float) and math.isclose(actual, expected, rel_tol=1e-12, abs_tol=1e-12)
    return type(actual) is type(expected) and actual == expected

def check_xlsx(filename):
    formulas = openpyxl.load_workbook(filename, data_only=False)
    values = openpyxl.load_workbook(filename, data_only=True)
    sample = INPUTS['sales']
    require(formulas.sheetnames == [s['name'] for s in sample['sheets']], 'XLSX sheet names/order changed')
    count = 0
    for sheet_index,sheet in enumerate(sample['sheets'],1):
        with ZipFile(filename) as archive:
            raw = read_xml(archive,f'xl/worksheets/sheet{sheet_index}.xml')
        ws, cached = formulas[sheet['name']], values[sheet['name']]
        require(ws.max_row == len(sheet['rows'])+1 and ws.max_column == len(sheet['columns']), 'XLSX table dimensions changed')
        require(ws.freeze_panes == 'A2', 'Frozen header lost')
        require(ws.auto_filter.ref == f'A1:{openpyxl.utils.get_column_letter(len(sheet["columns"]))}{len(sheet["rows"])+1}', 'Filter range changed')
        require(ws.print_title_rows in ('1:1', '$1:$1'), 'Repeated print header lost')
        require(ws.page_setup.fitToWidth == 1, 'Print width no longer fits page')
        for c, column in enumerate(sheet['columns'], 1):
            cell = ws.cell(1, c)
            require(cell.value == column['title'] and cell.font.bold, 'XLSX header content/bold changed')
            require(str(cell.fill.fgColor.rgb).endswith('264D73'), 'XLSX header fill changed')
        for r, row in enumerate(sheet['rows'], 2):
            for c, item in enumerate(row, 1):
                cell, value = ws.cell(r, c), cached.cell(r, c).value
                if isinstance(item, dict) and 'formula' in item:
                    require(cell.data_type == 'f' and formula_normal(cell.value) == formula_normal(item['formula']), f'Formula changed: {sheet["name"]}!{cell.coordinate}')
                    expected = ANSWERS[(sheet['name'], cell.coordinate)]
                    if expected == '' and value is None:
                        namespace={'s':'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}
                        node=raw.find(f".//s:c[@r='{cell.coordinate}']",namespace)
                        v=node.find('s:v',namespace)
                        require(node.get('t')=='str' and v is not None and not v.text,'Empty-string cache lost its type or value element')
                        value=''  # openpyxl collapses an explicitly typed empty <v/> to None
                    require(same_value(value, expected), f'Wrong typed result {sheet["name"]}!{cell.coordinate}: {value!r} vs {expected!r}')
                elif isinstance(item, dict) and 'date' in item:
                    require(isinstance(value, dt.datetime) and value.date().isoformat() == item['date'], 'Date type/value lost')
                    require('yyyy' in cell.number_format.lower() and 'mm' in cell.number_format.lower(), 'Date format lost')
                else:
                    require(same_value(value, item), f'Cell value changed: {sheet["name"]}!{cell.coordinate}')
                if sheet['columns'][c-1].get('format') == 'currency' and value is not None:
                    require('0.00' in cell.number_format, 'Two-decimal currency format lost')
                count += 1
    return f'{count} data cells, 10 typed formula results, dates, headers, filters, freeze and print settings match'

def pdf_pages(filename):
    data = subprocess.check_output(['pdftotext', '-bbox-layout', str(filename), '-'])
    root = ET.fromstring(data)
    pages = root.findall('.//h:page', NS)
    require(pages, 'No PDF pages')
    return pages

def page_text(page):
    words = sorted(page.findall('.//h:word', NS), key=lambda w: (float(w.attrib['yMin']), float(w.attrib['xMin'])))
    rows = []
    for word in words:
        y = float(word.attrib['yMin'])
        if not rows or abs(rows[-1][0]-y)>2:
            rows.append((y, [word]))
        else:
            rows[-1][1].append(word)
    return '\n'.join(' '.join(w.text or '' for w in sorted(items,key=lambda w:float(w.attrib['xMin']))) for _,items in rows)

def report_content(text):
    sample = INPUTS['report']
    for value in [sample['title'], sample['sections'][0]['heading'], *sample['sections'][0]['paragraphs']]:
        require(normal(value) in normal(text), 'Report paragraph missing from PDF')
    rows = re.findall(r'^\s*(\d+)\s+한글\s+표\s+내용\s*$', text, re.MULTILINE)
    require(rows == [str(i) for i in range(1,81)], f'Report PDF row sequence/count wrong: {len(rows)}')

def pdf_content(filename, sample, pages, reference=False):
    texts = [page_text(page) for page in pages]
    all_text = '\n'.join(texts)
    if sample in ('report', 'report_pdf'):
        report_content(all_text)
        for text in texts:
            require('항목' in text and '설명' in text, 'Report header missing on a page')
    elif sample == 'slides':
        require(len(pages) == 2, 'Slides must have exactly two pages')
        for i, text in enumerate(texts):
            for value in expected_pptx()[i]:
                require(normal(value) in normal(text), f'Slide {i+1} text missing')
    elif sample == 'sales':
        require(len(pages) == 2, 'Sample sheets must have exactly two PDF pages')
        rows = [['매출','날짜','항목','금액','세금','2026-09-07','상품 A','100.00','10.00','2026-09-08','상품 B','200.00','20.00','합계','300.00','30.00'],
                ['요약','항목','결과','총액','300','평균','150','달성','TRUE','판정','충족','문자 숫자','123','빈 문자열']]
        expected_rows = [['2026-09-07 상품 A 100.00 10.00','2026-09-08 상품 B 200.00 20.00','합계 300.00 30.00'],
                         ['총액 300','평균 150','달성 TRUE','판정 충족','문자 숫자 123','빈 문자열']]
        headings = [('매출','날짜 항목 금액 세금'),('요약','항목 결과')]
        for i,text in enumerate(texts):
            actual_rows = [normal(line) for line in text.splitlines() if line.strip() and normal(line) not in {normal(v) for v in headings[i]} and not re.fullmatch(r'\d+\s*/\s*\d+',line.strip())]
            if i == 1:
                for index,(label,expected) in enumerate([('총액',300),('평균',150)]):
                    match=re.fullmatch(label+r'([0-9,]+(?:\.[0-9]+)?)',actual_rows[index])
                    require(match is not None and same_value(float(match[1].replace(',','')),expected),'General-format numeric PDF value changed')
                    canonical=normal(expected_rows[i][index])
                    if actual_rows[index] != canonical:
                        RESULT.setdefault('presentation_differences',[]).append({'file':str(filename.relative_to(ROOT)),'page':2,'display':actual_rows[index],'canonical':canonical,'reason':'General number format; same numeric value'})
                    actual_rows[index]=canonical
            require(actual_rows == [normal(line) for line in expected_rows[i]], f'Sheet PDF row/cell order changed on page {i+1}: {actual_rows}')
        for i, items in enumerate(rows):
            for value in items:
                if reference and value in ('매출','요약'): continue  # sheet names are metadata, not printed cells
                require(normal(value) in normal(texts[i]), f'Sheet PDF missing {value!r}')

def geometry(pages):
    for p, page in enumerate(pages, 1):
        w, h = float(page.attrib['width']), float(page.attrib['height'])
        lines = []
        for word in page.findall('.//h:word', NS):
            x1,y1,x2,y2 = [float(word.attrib[key]) for key in ('xMin','yMin','xMax','yMax')]
            require(x1 >= 0 and y1 >= 0 and x2 <= w and y2 <= h, f'Text clipped on page {p}')
        for line in page.findall('.//h:line', NS):
            lines.append(tuple(float(line.attrib[key]) for key in ('xMin','yMin','xMax','yMax')))
        for i,a in enumerate(lines):
            for b in lines[i+1:]:
                overlap_x = min(a[2],b[2])-max(a[0],b[0])
                overlap_y = min(a[3],b[3])-max(a[1],b[1])
                require(not(overlap_x > 1 and overlap_y > 1), f'Text lines overlap on page {p}')

def render_and_contrast(filename, label, sample, reference):
    folder = ROOT/'review';folder.mkdir(exist_ok=True)
    prefix = folder/label
    for stale in folder.glob(label+'-*.png'): stale.unlink()
    subprocess.run(['pdftoppm','-scale-to','1100','-png',str(filename),str(prefix)],check=True,capture_output=True)
    images = sorted(folder.glob(label+'-*.png'))
    for file in images:
        image = Image.open(file).convert('RGB')
        require(image.getextrema() != ((255,255),(255,255),(255,255)), 'Blank PDF image')
        # Generated PDFs and formatted XLSX/PPTX references should retain the accent.
        if not reference or sample in ('sales','slides'):
            blue = [(i%image.width,i//image.width) for i,pixel in enumerate(image.getdata())
                    if max(abs(pixel[j]-[38,77,115][j]) for j in range(3)) <= 5]
            require(len(blue)>100, f'Accent/header color missing: {file.name}')
            if sample != 'slides':
                x1,y1 = min(p[0] for p in blue),min(p[1] for p in blue)
                x2,y2 = max(p[0] for p in blue),max(p[1] for p in blue)
                inner = image.crop((x1+3,y1+3,x2-3,y2-3))
                require(sum(min(p)>220 for p in inner.getdata())>10, f'Header text is not visible: {file.name}')
    return [str(p.relative_to(ROOT)) for p in images]

def check_pdf(filename, sample, reference=False):
    pages = pdf_pages(filename)
    pdf_content(filename,sample,pages,reference)
    geometry(pages)
    label = sample+('-reference' if reference else '-direct')
    images = render_and_contrast(filename,label,sample,reference)
    require(len(images)==len(pages),'Raster page count differs from PDF')
    RESULT['pdfs'].append({'name':label,'file':str(filename.relative_to(ROOT)),'pages':len(pages),'images':images})
    fonts = subprocess.check_output(['pdffonts',str(filename)],text=True)
    require('NotoSansCJK' in fonts, 'Expected Korean font missing')
    for line in fonts.splitlines()[2:]:
        if line.strip():
            require(re.search(r'\byes\s+(?:yes|no)\s+(?:yes|no)\s+\d+\s+\d+\s*$',line), 'PDF font is not embedded')
    return f'{len(pages)} pages: all sample text, bounds, line overlap, font embedding and raster contrast checked'

def negative_controls():
    def rewrite(source, target, entry, mutate):
        with ZipFile(source) as original, ZipFile(target, 'w') as changed:
            for info in original.infolist():
                data = original.read(info.filename)
                if info.filename == entry:
                    xml = ET.fromstring(data);mutate(xml);data = ET.tostring(xml,encoding='utf-8',xml_declaration=True)
                changed.writestr(info,data)
    def rejects(name, probe):
        try: probe()
        except AssertionError: RESULT['negative_controls'].append({'name':name,'status':'rejected'})
        else: raise AssertionError('Checker accepted corruption: '+name)
    with tempfile.TemporaryDirectory(prefix='office-negative-') as temp:
        folder = Path(temp)
        def remove_row(xml):
            table=xml.find('.//w:tbl',NS);table.remove(table.findall('w:tr',NS)[-1])
        file=folder/'row.docx';rewrite(ROOT/'report/document.docx',file,'word/document.xml',remove_row)
        rejects('missing DOCX table row',lambda:check_docx(file))
        def remove_bullet(xml):
            parents={child:parent for parent in xml.iter() for child in parent}
            for paragraph in xml.findall('.//a:p',NS):
                if 'DOCX' in ''.join(t.text or '' for t in paragraph.findall('.//a:t',NS)):
                    parents[paragraph].remove(paragraph);return
            raise AssertionError('Control paragraph not found')
        file=folder/'bullet.pptx';rewrite(ROOT/'slides/document.pptx',file,'ppt/slides/slide1.xml',remove_bullet)
        rejects('missing PPTX bullet',lambda:check_pptx(file))
        for name,entry,cell_id,kind,value in [('wrong cached result','sheet1.xml','C4','n','301'),('boolean changed to string','sheet2.xml','B4','str','TRUE'),('numeric string changed to number','sheet2.xml','B6','n','123'),('empty string changed to blank','sheet2.xml','B7','n','')]:
            def mutate(xml):
                namespace={'s':'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}
                cell=xml.find(f".//s:c[@r='{cell_id}']",namespace);cell.set('t',kind);cell.find('s:v',namespace).text=value
            file=folder/'changed.xlsx';rewrite(ROOT/'sales/document.xlsx',file,'xl/worksheets/'+entry,mutate)
            rejects(name,lambda:check_xlsx(file))
        def remove_filter(xml):
            namespace={'s':'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}
            node=xml.find('s:autoFilter',namespace);xml.remove(node)
        file=folder/'filter.xlsx';rewrite(ROOT/'sales/document.xlsx',file,'xl/worksheets/sheet1.xml',remove_filter)
        rejects('missing XLSX filter',lambda:check_xlsx(file))
        pages=pdf_pages(ROOT/'report/document.pdf')
        for page in pages:
            parents={child:parent for parent in page.iter() for child in parent}
            for word in page.findall('.//h:word',NS):
                if word.text == '80': parents[word].remove(word)
        rejects('missing PDF table row',lambda:pdf_content(None,'report',pages))
        pages=pdf_pages(ROOT/'slides/document.pdf');word=pages[0].find('.//h:word',NS);word.set('xMax',str(float(pages[0].attrib['width'])+10))
        rejects('clipped PDF text',lambda:geometry(pages))

def write_report():
    verified = set()
    for stem,item in INPUTS.items():
        ext=item['format'];verified.add(ROOT/stem/'document.pdf')
        if ext != 'pdf':
            verified.update([ROOT/stem/('document.'+ext),ROOT/stem/'reference'/('document.'+ext),ROOT/stem/'reference/document.pdf'])
        for file in ([stem+'.pdf'] if ext=='pdf' else [stem+'.'+ext,stem+'_view.pdf']):
            published=ROOT/file;source=ROOT/stem/('document'+published.suffix)
            require(published.read_bytes()==source.read_bytes(),'Published sample differs from verified source')
            verified.add(published)
    for path in sorted(verified):
        RESULT['artifacts'].append({'file':str(path.relative_to(ROOT)),'bytes':path.stat().st_size,'sha256':hashlib.sha256(path.read_bytes()).hexdigest()})
    reviewed=ROOT/'visual-review.json'
    RESULT['visual_review']={'status':'not_recorded'}
    if reviewed.exists():
        record=json.loads(reviewed.read_text())
        current={name:hashlib.sha256((ROOT/name).read_bytes()).hexdigest() for pdf in RESULT['pdfs'] for name in pdf['images']}
        RESULT['visual_review']={'status':'pass' if current==record['images'] and record.get('result')=='pass' else 'needs_review','reviewed_at':record['reviewed_at'],'pages':record['pages']}
    (ROOT/'verification.json').write_text(json.dumps(RESULT,ensure_ascii=False,indent=2))
    rows=''.join(f'<tr><td>{html.escape(c["name"])}</td><td>{c["status"]}</td><td>{html.escape(c["detail"])}</td></tr>' for c in RESULT['checks'])
    def page_image(filename):
        data='data:image/png;base64,'+base64.b64encode((ROOT/filename).read_bytes()).decode()
        return f'<a href="{data}" download="{html.escape(Path(filename).name)}"><img loading="lazy" src="{data}"></a>'
    pictures=''.join(f'<section><h2>{html.escape(p["name"])}</h2><p>{p["pages"]} pages · <a href="{html.escape(p["file"])}">PDF</a></p>'+''.join(page_image(img) for img in p['images'])+'</section>' for p in RESULT['pdfs'])
    differences=''.join('<li>'+html.escape(d['display']+' → '+d['canonical']+' ('+d['reason']+')')+'</li>' for d in RESULT.get('presentation_differences',[]))
    overview=f"<p>자동 검사 {len(RESULT['checks'])}건 · PDF {sum(p['pages'] for p in RESULT['pdfs'])}페이지 · 손상 대조군 {len(RESULT['negative_controls'])}건 감지 · 시각 검토 {RESULT['visual_review']['status']}</p><p>Office와 별도 생성 PDF의 배치는 다를 수 있습니다.</p><ul>{differences}</ul>"

    (ROOT/'verification.html').write_text('<!doctype html><html lang="ko"><meta charset="utf-8"><title>Office sample verification</title><style>body{font:15px sans-serif;margin:28px;background:#eef2f6;color:#18334d}table{border-collapse:collapse;background:white;width:100%}td{padding:10px;border:1px solid #ccd5df}section{margin:24px 0}img{width:240px;max-width:95%;vertical-align:top;margin:5px;border:1px solid #ccd5df}a{color:#264d73}</style><h1>Office sample verification: '+RESULT['status']+'</h1>'+overview+'<details><summary>검사 항목</summary><table>'+rows+'</table></details>'+pictures+'</html>')

try:
    for suffix in ['report/document.docx','report/reference/document.docx']:
        check(suffix,lambda suffix=suffix:check_docx(ROOT/suffix))
    for suffix in ['slides/document.pptx','slides/reference/document.pptx']:
        check(suffix,lambda suffix=suffix:check_pptx(ROOT/suffix))
    for suffix in ['sales/document.xlsx','sales/reference/document.xlsx']:
        check(suffix,lambda suffix=suffix:check_xlsx(ROOT/suffix))
    for stem in INPUTS:
        check(stem+' direct PDF',lambda stem=stem:check_pdf(ROOT/stem/'document.pdf',stem))
        if stem != 'report_pdf':
            check(stem+' LibreOffice PDF',lambda stem=stem:check_pdf(ROOT/stem/'reference/document.pdf',stem,True))
    negative_controls()
    RESULT['status']='pass'
except Exception as error:
    RESULT['status']='fail';RESULT['error']=str(error)
    raise
finally:
    write_report()
print(f"PASS: {len(RESULT['checks'])} checks; {sum(p['pages'] for p in RESULT['pdfs'])} PDF pages; {len(RESULT['negative_controls'])} corruption controls rejected")
